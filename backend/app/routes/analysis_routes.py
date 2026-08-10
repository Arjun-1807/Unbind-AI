import asyncio
import io
import json
import logging
from datetime import datetime, timezone

import httpx
from bson import ObjectId
from bson.errors import InvalidId
from fastapi import APIRouter, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from groq import APIConnectionError, APIStatusError, RateLimitError
from langsmith.run_helpers import tracing_context
from pymongo.errors import PyMongoError

from app.auth import get_current_user_id
from app.database import get_db
from app.schemas import (
    AnalyzeRequest,
    DocumentQuestionRequest,
    NegotiationDraftRequest,
    SimulateRequest,
)
from app.services import document_chat_service, reminder_service, vector_store
from app.services.analysis_service import ClauseExtractionError, analyze_contract
from app.services.negotiation_service import draft_negotiation_message
from app.services.ocr_service import OcrError, image_to_text, is_image_upload
from app.services.quota_service import ANALYSIS, QUERY, release, reserve

# Reject oversized image uploads before OCR. Phone photos are a few MB; anything
# much larger is likely not a document photo and wastes vision tokens/time.
_MAX_IMAGE_BYTES = 15 * 1024 * 1024

# Cap document uploads too. 25 MB comfortably covers a scanned several-hundred-
# page contract.
_MAX_DOCUMENT_BYTES = 25 * 1024 * 1024

# Block size for the capped upload read. Big enough that a legitimate 25 MB file
# is 25 awaits, small enough that the overshoot past the cap is negligible.
_UPLOAD_BLOCK_BYTES = 1024 * 1024

# Fixed, client-safe copy for failures whose real cause must not be echoed back.
# Exception text carries library internals, file paths and sometimes
# configuration ("GROQ_API_KEY is not set"); the client gets a stable code and
# the details go to the log.
_GENERIC_ERROR_CODE = "ANALYSIS_FAILED"
_GENERIC_ERROR_DETAIL = "The analysis could not be completed. Please try again."
_PDF_ERROR_DETAIL = "This PDF could not be read. Try re-exporting it, or upload it as a DOCX."
_DOCX_ERROR_DETAIL = "This DOCX could not be read. Try re-exporting it, or upload it as a PDF."

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/analysis", tags=["analysis"])


class PreflightRejection(HTTPException):
    """A 4xx raised before any LLM call was made.

    Exists purely so the refund rule can be decided structurally instead of by
    guessing from an exception's type or message: an upload we rejected while
    reading or parsing it cost CPU but not a single token, so handing the quota
    unit back is free. The never-refunded attempt counter in ``quota_service``
    is what stops someone looping on these.
    """


# Upstream/our-side failures where the completion never happened, so nothing was
# billed. An allowlist, because guessing wrong in the other direction is what
# costs money.
_INFRASTRUCTURE_ERRORS: tuple[type[BaseException], ...] = (
    APIConnectionError,  # also covers APITimeoutError
    RateLimitError,  # upstream refused the request outright, nothing generated
    httpx.TransportError,  # connect/read/write timeouts and dropped sockets
    PyMongoError,  # we produced a result and then failed to store it
    ConnectionError,
    TimeoutError,  # asyncio.TimeoutError is an alias of this on 3.11+
)


def _is_infrastructure_failure(exc: BaseException) -> bool:
    """True when the failure is ours or an upstream's, not the request's."""
    if isinstance(exc, PreflightRejection):
        return False
    if isinstance(exc, HTTPException):
        # Our own 4xx are rejections of the request; only a 5xx is our fault.
        return exc.status_code >= 500
    if isinstance(exc, APIStatusError):
        return exc.status_code >= 500
    return isinstance(exc, _INFRASTRUCTURE_ERRORS)


def _should_refund(exc: BaseException) -> bool:
    """True when this failure cost us no LLM tokens, so the quota unit is free to give back.

    The single question this asks is "did we already pay Groq for this?".
    Refundable: an infrastructure failure (upstream 5xx, timeout, dropped
    connection, database outage), an upload rejected before any model call
    (:class:`PreflightRejection`), and an image that never reached the vision
    model (:class:`OcrError` — HEIC, undecodable bytes, a decompression bomb;
    every one of those is raised before ``ocr_complete``).

    Not refundable, and this is the whole point: ``OCR_INSUFFICIENT_TEXT`` (the
    vision completion happened and returned nothing useful),
    ``NOT_A_LEGAL_DOCUMENT`` (a classifier completion reached that verdict) and
    :class:`ClauseExtractionError` (up to a few hundred chunk completions, all
    billed). Refunding those is what made a failure loop free and let one free
    account drive unbounded spend — the counter went back where it started, so
    the gate never closed.

    Anything unrecognised (a bare ``RuntimeError``, say) falls on the
    no-refund side deliberately: being wrong there costs one user one unit of
    their daily allowance, versus an unmetered LLM bill.
    """
    return isinstance(exc, (PreflightRejection, OcrError)) or _is_infrastructure_failure(exc)


async def _release_if_refundable(
    exc: BaseException, user_id: str, counter=ANALYSIS, *, reserved: bool = True
) -> None:
    """Refund the reserved quota unit only if ``exc`` cost us no LLM tokens."""
    if _should_refund(exc):
        await release(user_id, counter, reserved=reserved)


async def _schedule_reminders(analysis_id: str, user_id: str, file_name: str, result: dict) -> None:
    """Derive deadline reminders from a finished analysis.

    Wrapped so a reminder failure can never fail the analysis the user waited
    for — reminders are a bonus on top of the result, not part of it.
    """
    await reminder_service.generate_safely(
        analysis_id, user_id, file_name, result.get("keyDates", []) or []
    )


def _object_id(raw: str, what: str = "analysis") -> ObjectId:
    """Parse a path id, answering 400 rather than 500 on a malformed one."""
    try:
        return ObjectId(raw)
    except (InvalidId, TypeError) as e:
        raise HTTPException(status_code=400, detail=f"Invalid {what} ID") from e


def _utcnow_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _size_limit(is_image: bool) -> tuple[int, str]:
    """The byte cap and the client-facing code for an upload of this kind."""
    if is_image:
        return _MAX_IMAGE_BYTES, "IMAGE_TOO_LARGE"
    return _MAX_DOCUMENT_BYTES, "FILE_TOO_LARGE"


async def _read_upload_within_limit(request: Request, file: UploadFile, is_image: bool) -> bytes:
    """Read an upload into memory, aborting the moment it exceeds its cap.

    The cap used to be applied to ``await file.read()`` — i.e. to the length of
    bytes that were already, in their entirety, in RAM — which enforced nothing
    at all. Reading in blocks and checking a running total means the process
    never holds more than the limit plus one block, so the 413 is real.

    ``Content-Length`` is consulted first as a cheap short-circuit. It counts
    multipart framing as well as the file, so it can only ever over-estimate:
    safe to reject on, never safe to trust as a substitute for counting.
    """
    limit, code = _size_limit(is_image)

    declared = request.headers.get("content-length")
    if declared and declared.strip().isdigit() and int(declared) > limit:
        raise PreflightRejection(status_code=413, detail=code)

    blocks: list[bytes] = []
    total = 0
    while True:
        block = await file.read(_UPLOAD_BLOCK_BYTES)
        if not block:
            break
        total += len(block)
        if total > limit:
            # Stop before this block joins the others, so the peak is bounded.
            raise PreflightRejection(status_code=413, detail=code)
        blocks.append(block)
    return b"".join(blocks)


async def _read_upload(request: Request, file: UploadFile) -> tuple[bytes, str, bool]:
    """Resolve an upload to ``(content, file_name, is_image)`` under its size cap."""
    file_name = file.filename or "document"
    is_image = is_image_upload(file.content_type, file_name)
    content = await _read_upload_within_limit(request, file, is_image)
    return content, file_name, is_image


def _sse_event(event: str, data: dict) -> str:
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"


async def _stream_analysis(
    text: str | None,
    role: str,
    file_name: str,
    user_id: str,
    tracing_tags: list[str],
    ocr_source: tuple[bytes, str | None, str] | None = None,
    reserved: bool = False,
):
    """Run analyze_contract while streaming progress events over SSE.

    The pipeline itself is a single awaited coroutine, so progress updates
    (fired synchronously from inside it via on_progress) are relayed to the
    client through a queue drained by a concurrently running consumer task.

    When ``ocr_source`` (raw image bytes, content-type, filename) is given, the
    image is OCR'd here first — emitting a progress event so the client sees
    feedback during the vision call — and the transcription becomes ``text``.

    ``reserved`` says whether the caller consumed a unit of the daily quota. It
    is handed back only for a failure that cost us no LLM tokens (see
    :func:`_should_refund`). Every outcome that a completion was already billed
    for keeps the unit — refunding those is what made "upload a photo of a blank
    wall in a loop" a free way to spend someone else's Groq budget.
    """
    if ocr_source is not None:
        yield _sse_event("progress", {"stage": "ocr", "message": "Reading text from your image…"})
        try:
            text = await image_to_text(*ocr_source)
        except OcrError as e:
            # The user's own image: HEIC, undecodable, or a decompression bomb —
            # all raised before the vision call, so the unit is refundable.
            # OcrError messages are written for humans, so passing them through
            # leaks nothing.
            await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
            yield _sse_event("error", {"code": e.code, "detail": e.message})
            return
        except Exception as e:
            await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
            logger.exception("Vision OCR failed for user %s", user_id)
            yield _sse_event(
                "error", {"code": _GENERIC_ERROR_CODE, "detail": _GENERIC_ERROR_DETAIL}
            )
            return
        if not text or len(text.strip()) < 50:
            # The vision completion was made and billed; no refund.
            yield _sse_event("error", {"code": "OCR_INSUFFICIENT_TEXT"})
            return

    queue: asyncio.Queue = asyncio.Queue()

    def on_progress(stage: str, detail: dict) -> None:
        queue.put_nowait({"stage": stage, **detail})

    async def run_pipeline() -> dict:
        with tracing_context(
            metadata={
                "endpoint": "analysis.analyze.stream",
                "user_id": user_id,
                "file_name": file_name,
                "role": role,
                "text_length": len(text or ""),
            },
            tags=tracing_tags,
        ):
            return await analyze_contract(text, role, user_id=user_id, on_progress=on_progress)

    pipeline_task = asyncio.create_task(run_pipeline())

    try:
        while not pipeline_task.done():
            try:
                event = await asyncio.wait_for(queue.get(), timeout=0.5)
                yield _sse_event("progress", event)
            except asyncio.TimeoutError:
                continue

        # Drain any events emitted right before completion.
        while not queue.empty():
            yield _sse_event("progress", queue.get_nowait())

        try:
            result = pipeline_task.result()
        except ValueError as e:
            if str(e) == "NOT_A_LEGAL_DOCUMENT":
                # A classifier completion was spent reaching this verdict.
                yield _sse_event("error", {"code": "NOT_A_LEGAL_DOCUMENT"})
                return
            logger.exception("Streaming analysis failed for user %s", user_id)
            yield _sse_event(
                "error", {"code": _GENERIC_ERROR_CODE, "detail": _GENERIC_ERROR_DETAIL}
            )
            return
        except ClauseExtractionError as e:
            # Every chunk completion was made and billed before we got here, so
            # this is emphatically not refundable. The message is ours, written
            # for the user, so it is safe to pass on.
            logger.warning("Clause extraction produced nothing for user %s: %s", user_id, e)
            yield _sse_event("error", {"code": "NO_CLAUSES_EXTRACTED", "detail": str(e)})
            return
        except Exception as e:
            # Catch-all, not just RuntimeError: letting an unexpected exception
            # escape the generator truncates the SSE stream, and the client sees
            # a dead connection instead of an error it can render.
            await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
            logger.exception("Streaming analysis failed for user %s", user_id)
            yield _sse_event(
                "error", {"code": _GENERIC_ERROR_CODE, "detail": _GENERIC_ERROR_DETAIL}
            )
            return

        db = get_db()
        doc = {
            "userId": user_id,
            "fileName": file_name,
            "analysisDate": _utcnow_iso(),
            "analysisResult": result,
            "documentText": text,
        }
        inserted = await db.analyses.insert_one(doc)
        await _schedule_reminders(str(inserted.inserted_id), user_id, file_name, result)

        yield _sse_event(
            "result",
            {
                "id": str(inserted.inserted_id),
                "userId": user_id,
                "fileName": file_name,
                "analysisDate": doc["analysisDate"],
                "analysisResult": result,
                "documentText": text,
            },
        )
    finally:
        if not pipeline_task.done():
            pipeline_task.cancel()


@router.post("/analyze")
async def analyze(body: AnalyzeRequest, request: Request):
    """Analyse contract text and return the full result."""
    user_id = await get_current_user_id(request)
    reserved = await reserve(user_id, ANALYSIS)

    try:
        try:
            with tracing_context(
                metadata={
                    "endpoint": "analysis.analyze",
                    "user_id": user_id,
                    "file_name": body.fileName,
                    "role": body.role,
                    "text_length": len(body.text or ""),
                },
                tags=["analysis", "api", "text"],
            ):
                result = await analyze_contract(body.text, body.role, user_id=user_id)
        except ValueError as e:
            if str(e) == "NOT_A_LEGAL_DOCUMENT":
                raise HTTPException(status_code=422, detail="NOT_A_LEGAL_DOCUMENT") from e
            raise
        except ClauseExtractionError as e:
            raise HTTPException(status_code=422, detail=str(e)) from e
    except HTTPException as e:
        # Both rejections above are verdicts a completion was already billed for,
        # so nothing is refunded; routed through the same helper so this stays
        # true if the block grows.
        await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
        raise
    except Exception as e:
        # Refund only what we broke. A rejected document or an unparseable model
        # response has already been paid for in tokens, and handing the quota
        # back for it is a free retry loop.
        await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
        logger.exception("Analysis failed for user %s", user_id)
        raise HTTPException(
            status_code=502 if _is_infrastructure_failure(e) else 500,
            detail=_GENERIC_ERROR_DETAIL,
        ) from e

    # Persist to DB
    db = get_db()
    doc = {
        "userId": user_id,
        "fileName": body.fileName,
        "analysisDate": _utcnow_iso(),
        "analysisResult": result,
        "documentText": body.text,
    }
    inserted = await db.analyses.insert_one(doc)
    await _schedule_reminders(str(inserted.inserted_id), user_id, body.fileName, result)

    return {
        "id": str(inserted.inserted_id),
        "userId": user_id,
        "fileName": body.fileName,
        "analysisDate": doc["analysisDate"],
        "analysisResult": result,
        "documentText": body.text,
    }


@router.post("/upload")
async def upload_and_analyze(
    request: Request,
    file: UploadFile = File(...),
    role: str = Form(""),
):
    """Upload a file (PDF or text), extract text, and analyse."""
    user_id = await get_current_user_id(request)
    # Size-check and slurp the body *before* reserving: an oversized upload costs
    # no LLM work, so it should cost no quota either, and this way the 413 needs
    # no refund at all.
    content, file_name, is_image = await _read_upload(request, file)
    reserved = await reserve(user_id, ANALYSIS)

    try:
        # Determine file type and extract text (images go through vision OCR).
        if is_image:
            try:
                text = await image_to_text(content, file.content_type, file_name)
            except OcrError as e:
                # Every OcrError is raised before the vision call, so this is
                # refundable — matching the streaming endpoint's behaviour.
                raise PreflightRejection(status_code=422, detail=e.code) from e
            if not text or len(text.strip()) < 50:
                # The vision completion ran and was billed, so unlike the
                # rejections above this one is not refundable.
                raise HTTPException(status_code=422, detail="OCR_INSUFFICIENT_TEXT")
        else:
            text = await _extract_text_from_upload(content, file.content_type, file_name)
            if not text or len(text.strip()) < 50:
                raise PreflightRejection(
                    status_code=422,
                    detail="Not enough text extracted from the document.",
                )

        try:
            with tracing_context(
                metadata={
                    "endpoint": "analysis.upload",
                    "user_id": user_id,
                    "file_name": file_name,
                    "file_content_type": file.content_type or "",
                    "role": role,
                    "text_length": len(text or ""),
                },
                tags=["analysis", "api", "upload"],
            ):
                result = await analyze_contract(text, role, user_id=user_id)
        except ValueError as e:
            if str(e) == "NOT_A_LEGAL_DOCUMENT":
                raise HTTPException(status_code=422, detail="NOT_A_LEGAL_DOCUMENT") from e
            raise
        except ClauseExtractionError as e:
            raise HTTPException(status_code=422, detail=str(e)) from e
    except HTTPException as e:
        # A 4xx says the upload was at fault. Refundable only when we hadn't paid
        # for a completion yet (PreflightRejection); OCR_INSUFFICIENT_TEXT and
        # NOT_A_LEGAL_DOCUMENT had, so they stand.
        await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
        raise
    except Exception as e:
        await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
        logger.exception("Upload analysis failed for user %s", user_id)
        raise HTTPException(
            status_code=502 if _is_infrastructure_failure(e) else 500,
            detail=_GENERIC_ERROR_DETAIL,
        ) from e

    # Persist
    db = get_db()
    doc = {
        "userId": user_id,
        "fileName": file_name,
        "analysisDate": _utcnow_iso(),
        "analysisResult": result,
        "documentText": text,
    }
    inserted = await db.analyses.insert_one(doc)
    await _schedule_reminders(str(inserted.inserted_id), user_id, file_name, result)

    return {
        "id": str(inserted.inserted_id),
        "userId": user_id,
        "fileName": file_name,
        "analysisDate": doc["analysisDate"],
        "analysisResult": result,
        "documentText": text,
    }


@router.post("/analyze/stream")
async def analyze_stream(body: AnalyzeRequest, request: Request):
    """Analyse contract text, streaming per-clause progress via SSE."""
    user_id = await get_current_user_id(request)
    reserved = await reserve(user_id, ANALYSIS)

    return StreamingResponse(
        _stream_analysis(
            body.text,
            body.role,
            body.fileName,
            user_id,
            ["analysis", "api", "text", "stream"],
            reserved=reserved,
        ),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.post("/upload/stream")
async def upload_and_analyze_stream(
    request: Request,
    file: UploadFile = File(...),
    role: str = Form(""),
):
    """Upload a file (PDF or text), extract text, and stream analysis progress via SSE."""
    user_id = await get_current_user_id(request)
    # Same ordering as /upload: cap the body first (no quota spent on a 413),
    # then reserve.
    content, file_name, is_image = await _read_upload(request, file)
    reserved = await reserve(user_id, ANALYSIS)

    try:
        # Image uploads: OCR happens inside the stream (with a progress event) so
        # the client gets feedback during the vision call. Pass raw bytes through.
        if is_image:
            return StreamingResponse(
                _stream_analysis(
                    None,
                    role,
                    file_name,
                    user_id,
                    ["analysis", "api", "upload", "stream", "ocr"],
                    ocr_source=(content, file.content_type, file_name),
                    reserved=reserved,
                ),
                media_type="text/event-stream",
                headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
            )

        text = await _extract_text_from_upload(content, file.content_type, file_name)

        if not text or len(text.strip()) < 50:
            raise PreflightRejection(
                status_code=422,
                detail="Not enough text extracted from the document.",
            )
    except HTTPException as e:
        await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
        raise
    except Exception as e:
        await _release_if_refundable(e, user_id, ANALYSIS, reserved=reserved)
        logger.exception("Upload text extraction failed for user %s", user_id)
        raise HTTPException(
            status_code=502 if _is_infrastructure_failure(e) else 500,
            detail=_GENERIC_ERROR_DETAIL,
        ) from e

    return StreamingResponse(
        _stream_analysis(
            text,
            role,
            file_name,
            user_id,
            ["analysis", "api", "upload", "stream"],
            reserved=reserved,
        ),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.get("/history")
async def history(
    request: Request,
    limit: int = Query(50, ge=1, le=200),
    skip: int = Query(0, ge=0),
):
    """Return a page of the user's analyses, newest first.

    ``documentText`` is deliberately omitted: it is by far the largest field and
    the list view never renders it. Clients fetch the full record from
    ``/history/{id}`` when the user actually opens one.
    """
    user_id = await get_current_user_id(request)
    db = get_db()
    cursor = (
        db.analyses.find({"userId": user_id}, {"documentText": 0})
        .sort("analysisDate", -1)
        .skip(skip)
        .limit(limit)
    )
    results = []
    async for doc in cursor:
        results.append(
            {
                "id": str(doc["_id"]),
                "userId": doc["userId"],
                "fileName": doc["fileName"],
                "analysisDate": doc["analysisDate"],
                "analysisResult": doc["analysisResult"],
            }
        )
    return results


@router.get("/history/{analysis_id}")
async def get_analysis(analysis_id: str, request: Request):
    """Return a single analysis by id, including its full document text."""
    user_id = await get_current_user_id(request)
    db = get_db()

    doc = await db.analyses.find_one({"_id": _object_id(analysis_id), "userId": user_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Analysis not found")
    return {
        "id": str(doc["_id"]),
        "userId": doc["userId"],
        "fileName": doc["fileName"],
        "analysisDate": doc["analysisDate"],
        "analysisResult": doc["analysisResult"],
        "documentText": doc.get("documentText", ""),
    }


@router.delete("/history/{analysis_id}")
async def delete_analysis(analysis_id: str, request: Request):
    """Delete a single analysis by id, plus everything derived from it."""
    user_id = await get_current_user_id(request)
    db = get_db()

    result = await db.analyses.delete_one({"_id": _object_id(analysis_id), "userId": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Analysis not found")

    # Deleting a document must delete what we derived from it — its embeddings
    # and its Q&A history — or the user's "delete" leaves their contract text
    # behind in chunk form.
    await vector_store.delete_index(analysis_id, user_id)
    await reminder_service.delete_for_analysis(analysis_id, user_id)
    try:
        await document_chat_service.clear_conversation(analysis_id, user_id)
    except Exception:
        logger.exception("Failed to clear chat for deleted analysis %s", analysis_id)

    return {"ok": True}


# ──── Document Q&A ────


async def _load_owned_analysis(analysis_id: str, user_id: str) -> dict:
    """Fetch an analysis the caller owns, or 404. Never leaks another user's doc."""
    db = get_db()
    doc = await db.analyses.find_one({"_id": _object_id(analysis_id), "userId": user_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Analysis not found")
    return doc


@router.post("/{analysis_id}/chat")
async def ask_document(analysis_id: str, body: DocumentQuestionRequest, request: Request):
    """Answer a question about an analysed document, with citations."""
    user_id = await get_current_user_id(request)
    doc = await _load_owned_analysis(analysis_id, user_id)

    reserved = await reserve(user_id, QUERY)
    try:
        with tracing_context(
            metadata={
                "endpoint": "analysis.chat",
                "user_id": user_id,
                "analysis_id": analysis_id,
                "question": body.question,
                "question_length": len(body.question),
            },
            tags=["analysis", "api", "chat"],
        ):
            return await document_chat_service.answer_question(
                analysis_id,
                user_id,
                doc.get("documentText", ""),
                body.question,
            )
    except Exception as e:
        await _release_if_refundable(e, user_id, QUERY, reserved=reserved)
        raise


@router.get("/{analysis_id}/chat")
async def get_document_chat(analysis_id: str, request: Request):
    """Return the stored Q&A conversation for an analysis."""
    user_id = await get_current_user_id(request)
    await _load_owned_analysis(analysis_id, user_id)
    messages = await document_chat_service.get_conversation(analysis_id, user_id)
    return [
        {
            "role": m.get("role"),
            "content": m.get("content", ""),
            "citations": m.get("citations", []),
            "createdAt": (
                m["createdAt"].isoformat()
                if hasattr(m.get("createdAt"), "isoformat")
                else m.get("createdAt")
            ),
        }
        for m in messages
    ]


@router.delete("/{analysis_id}/chat")
async def clear_document_chat(analysis_id: str, request: Request):
    """Clear the Q&A conversation for an analysis."""
    user_id = await get_current_user_id(request)
    await _load_owned_analysis(analysis_id, user_id)
    await document_chat_service.clear_conversation(analysis_id, user_id)
    return {"ok": True}


@router.post("/simulate")
async def simulate(body: SimulateRequest, request: Request):
    """Answer a one-off question about contract text, with citations.

    Retained for the published CLI, which posts raw document text with no
    analysis id and reads back ``result``. The web app uses
    ``POST /{analysis_id}/chat`` instead, which shares this exact answering path
    but adds a stored conversation.
    """
    user_id = await get_current_user_id(request)
    # An analysisId keys a persisted vector index (up to ~800 chunk records), so
    # an unchecked one lets a caller mint permanent index rows under ids they
    # have nothing to do with. Every vector-store read is already scoped by
    # userId, so this is not a cross-user read — but the write side needs the
    # same ownership rule the /{analysis_id}/chat path already applies.
    if body.analysisId:
        await _load_owned_analysis(body.analysisId, user_id)

    # Metered: each call is a HyDE completion plus retrieval, so leaving it
    # unbounded is an unbounded bill.
    reserved = await reserve(user_id, QUERY)
    try:
        with tracing_context(
            metadata={
                "endpoint": "analysis.simulate",
                "user_id": user_id,
                "scenario": body.scenario,
                "scenario_length": len(body.scenario or ""),
                "document_length": len(body.documentText or ""),
                "analysis_id": body.analysisId,
            },
            tags=["analysis", "api", "simulate"],
        ):
            result = await document_chat_service.answer_standalone(
                body.documentText,
                body.scenario,
                user_id=user_id,
                analysis_id=body.analysisId,
            )
    except Exception as e:
        await _release_if_refundable(e, user_id, QUERY, reserved=reserved)
        raise
    # ``result`` is the legacy field name the CLI reads; ``citations`` lets newer
    # clients link answer markers back to the document.
    return {"result": result["answer"], "citations": result["citations"]}


@router.post("/negotiation-message")
async def negotiation_message(body: NegotiationDraftRequest, request: Request):
    """Draft a ready-to-send negotiation message from selected clause changes."""
    user_id = await get_current_user_id(request)
    reserved = await reserve(user_id, QUERY)
    try:
        with tracing_context(
            metadata={
                "endpoint": "analysis.negotiation_message",
                "user_id": user_id,
                "point_count": len(body.points),
                "tone": body.tone,
                "format": body.format,
            },
            tags=["analysis", "api", "negotiation"],
        ):
            return await draft_negotiation_message(body)
    except Exception as e:
        await _release_if_refundable(e, user_id, QUERY, reserved=reserved)
        raise


# ──── File text extraction helpers ────


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


async def _extract_text_from_upload(
    content: bytes, content_type: str | None, file_name: str
) -> str:
    """Dispatch to the right extractor based on content-type/filename, or decode as plain text.

    pdfplumber and python-docx are synchronous and CPU-bound — a large PDF parsed
    inline would stall the event loop for every other request on the worker — so
    they run in a thread, matching ocr_service's handling of Pillow.
    """
    if content_type == "application/pdf" or file_name.lower().endswith(".pdf"):
        return await asyncio.to_thread(_extract_pdf_text, content)
    if content_type == DOCX_CONTENT_TYPE or file_name.lower().endswith(".docx"):
        return await asyncio.to_thread(_extract_docx_text, content)
    return content.decode("utf-8", errors="replace")


def _extract_pdf_text(pdf_bytes: bytes) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            pages = []
            for page in pdf.pages:
                text = page.extract_text() or ""
                pages.append(text)
            return "\n\n".join(pages)
    except Exception:
        # Fallback to pypdf (the maintained successor to the archived PyPDF2;
        # PdfReader's API is unchanged).
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(pdf_bytes))
            pages = []
            for page in reader.pages:
                text = page.extract_text() or ""
                pages.append(text)
            return "\n\n".join(pages)
        except Exception as e:
            # The exception text is a parser's opinion of attacker-controlled
            # bytes — library internals, sometimes file paths. Log it, tell the
            # client something it can act on.
            logger.exception("PDF text extraction failed")
            raise PreflightRejection(status_code=422, detail=_PDF_ERROR_DETAIL) from e


def _extract_docx_text(docx_bytes: bytes) -> str:
    try:
        import docx

        document = docx.Document(io.BytesIO(docx_bytes))

        paragraphs = [p.text for p in document.paragraphs if p.text and p.text.strip()]

        table_text: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    table_text.append("\t".join(cells))

        return "\n".join(paragraphs + table_text)
    except Exception as e:
        logger.exception("DOCX text extraction failed")
        raise PreflightRejection(status_code=422, detail=_DOCX_ERROR_DETAIL) from e
